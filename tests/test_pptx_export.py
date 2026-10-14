"""PPTX 导出（每页一张幻灯片）的结构检查。

为什么值得单独测：PPTX 是"看起来能生成就行"的格式，但一个结构错误的文件
PowerPoint 会直接拒开或显示空白页，而 python-pptx 自己读回来往往还是正常的。
所以这里检查的是**结构性事实**：幻灯片数、幻灯片尺寸、每页图片的位置与缩放，
以及混排纸张时有没有被裁掉。

不校验渲染画面是否好看——那要靠人眼；这里只管"文件结构对不对"。
"""

import json
import os
import subprocess
import sys
import tempfile
import unittest

import pymupdf
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
PDF_EXE = os.path.join(REPO, "build", "bin", "pdf.exe")
EMU_PER_PT = 12700.0


def make_pdf(path, sizes, text="页"):
    doc = pymupdf.open()
    for i, (w, h) in enumerate(sizes):
        page = doc.new_page(width=w, height=h)
        page.insert_text((40, 60), f"{text}{i + 1}", fontsize=20)
    doc.save(path)
    doc.close()


def export(plan_path, out_path, fmt="pptx", dpi=100):
    p = subprocess.run(
        [PDF_EXE, "ws-export", "--format", fmt, "--output", out_path, "--dpi", str(dpi), plan_path],
        capture_output=True, text=True, errors="replace")
    return p.returncode, (p.stdout or "") + (p.stderr or "")


class TestPptxExport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not os.path.exists(PDF_EXE):
            raise unittest.SkipTest(f"找不到 {PDF_EXE}，先跑 scripts/build.ps1")
        cls.tmp = tempfile.mkdtemp(prefix="pdfguru-pptx-test-")

    def _plan(self, name, sizes):
        pdf = os.path.join(self.tmp, name + ".pdf")
        make_pdf(pdf, sizes)
        plan = os.path.join(self.tmp, name + ".json")
        with open(plan, "w", encoding="utf-8") as f:
            json.dump({"pages": [{"path": pdf, "index": i, "rotation": 0}
                                 for i in range(len(sizes))], "options": {}}, f)
        return plan

    def test_slide_count_and_size_match_pages(self):
        plan = self._plan("a4", [(595, 842)] * 3)
        out = os.path.join(self.tmp, "a4.pptx")
        code, log = export(plan, out, dpi=100)
        self.assertEqual(code, 0, log)
        self.assertTrue(os.path.exists(out), "没有产出 pptx")

        prs = Presentation(out)
        self.assertEqual(len(prs.slides), 3, "幻灯片数应等于页数")
        self.assertAlmostEqual(prs.slide_width / EMU_PER_PT, 595, delta=1)
        self.assertAlmostEqual(prs.slide_height / EMU_PER_PT, 842, delta=1)
        for slide in prs.slides:
            pics = [s for s in slide.shapes if s.shape_type == 13]
            self.assertEqual(len(pics), 1, "每张幻灯片应恰好一张页面图片")
            pic = pics[0]
            # 单页尺寸与幻灯片一致 -> 应当正好铺满，不留边
            self.assertAlmostEqual(pic.left, 0, delta=EMU_PER_PT)
            self.assertAlmostEqual(pic.top, 0, delta=EMU_PER_PT)
            self.assertAlmostEqual(pic.width / EMU_PER_PT, 595, delta=1)
            self.assertAlmostEqual(pic.height / EMU_PER_PT, 842, delta=1)

    def test_mixed_sizes_are_fitted_and_centred_not_cropped(self):
        # A2 横 + A4 竖：幻灯片取最大页，A4 竖页应等比缩放并水平居中
        plan = self._plan("mixed", [(1684, 1191), (595, 842)])
        out = os.path.join(self.tmp, "mixed.pptx")
        code, log = export(plan, out, dpi=80)
        self.assertEqual(code, 0, log)

        prs = Presentation(out)
        self.assertAlmostEqual(prs.slide_width / EMU_PER_PT, 1684, delta=1)
        self.assertAlmostEqual(prs.slide_height / EMU_PER_PT, 1191, delta=1)

        a2 = [s for s in prs.slides[0].shapes if s.shape_type == 13][0]
        self.assertAlmostEqual(a2.width / EMU_PER_PT, 1684, delta=1)
        self.assertAlmostEqual(a2.height / EMU_PER_PT, 1191, delta=1)

        a4 = [s for s in prs.slides[1].shapes if s.shape_type == 13][0]
        w, h = a4.width / EMU_PER_PT, a4.height / EMU_PER_PT
        # 长宽比必须保持，且完整落在幻灯片内（不能被裁）
        self.assertAlmostEqual(w / h, 595 / 842, delta=0.01)
        self.assertLessEqual(w, 1684 + 1)
        self.assertLessEqual(h, 1191 + 1)
        self.assertAlmostEqual(a4.left / EMU_PER_PT, (1684 - w) / 2, delta=1)
        # 竖直方向刚好占满
        self.assertAlmostEqual(h, 1191, delta=1)

    def test_oversized_page_is_clamped_within_powerpoint_limit(self):
        # PowerPoint 的幻灯片上限是 56 英寸 = 4032pt，超了它直接打不开文件
        plan = self._plan("huge", [(5000, 3000)])
        out = os.path.join(self.tmp, "huge.pptx")
        code, log = export(plan, out, dpi=40)
        self.assertEqual(code, 0, log)
        prs = Presentation(out)
        self.assertLessEqual(prs.slide_width / EMU_PER_PT, 4032 + 1)
        self.assertLessEqual(prs.slide_height / EMU_PER_PT, 4032 + 1)
        # 长宽比仍然保持
        self.assertAlmostEqual(
            (prs.slide_width / EMU_PER_PT) / (prs.slide_height / EMU_PER_PT),
            5000 / 3000, delta=0.02)

    def test_page_range_limits_slides(self):
        plan = self._plan("six", [(595, 842)] * 6)
        # 只取清单的前 2 页：模拟"仅导出选中页"
        with open(plan, "r", encoding="utf-8") as f:
            data = json.load(f)
        data["pages"] = data["pages"][:2]
        plan2 = os.path.join(self.tmp, "six2.json")
        with open(plan2, "w", encoding="utf-8") as f:
            json.dump(data, f)
        out = os.path.join(self.tmp, "six.pptx")
        code, log = export(plan2, out, dpi=60)
        self.assertEqual(code, 0, log)
        self.assertEqual(len(Presentation(out).slides), 2)


if __name__ == "__main__":
    unittest.main(verbosity=2)
